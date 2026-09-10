import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import sys
import os
import time
import concurrent.futures
import threading

exit_requested = threading.Event()

def input_listener():
    while not exit_requested.is_set():
        try:
            line = sys.stdin.readline()
            if line and line.strip().lower() == 'x':
                print("\n[INFO] 'x' entered by user. Terminating program...")
                os._exit(0)
        except Exception:
            break

def check_name_column(col_info, sheet):
    col_name, col_idx = col_info
    if col_name and str(col_name).endswith("Name"):
        unique_values = set()
        for row_idx in range(2, sheet.max_row + 1):
            cell_value = sheet.cell(row=row_idx, column=col_idx).value
            if cell_value is not None and str(cell_value).strip().lower() != "null":
                unique_values.add(str(cell_value).strip())
        if len(unique_values) == 1:
            return col_idx, col_name, unique_values.pop()
    return None

def check_null_column(col_idx, sheet, columns_to_delete):
    if col_idx in columns_to_delete:
        return None
    is_null_column = True
    for row_idx in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(row=row_idx, column=col_idx).value
        if cell_value is not None and str(cell_value).strip().lower() != "null" and str(cell_value).strip() != "":
            is_null_column = False
            break
    if is_null_column:
        return col_idx
    return None

def remove_null_columns(filename):
    total_start = time.time()
    
    ext = os.path.splitext(filename)[1].lower()
    
    # --- CHECKPOINT 1: Loading File ---
    t0 = time.time()
    try:
        if ext == '.csv':
            wb = openpyxl.Workbook()
            sheet = wb.active
            with open(filename, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                for r_idx, row in enumerate(reader, 1):
                    for c_idx, val in enumerate(row, 1):
                        sheet.cell(row=r_idx, column=c_idx, value=val)
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(filename)
            sheet = wb.active
        else:
            print("Error: Unsupported file type.")
            return
    except Exception as e:
        print(f"Error loading file: {e}")
        return
    print(f"[DEBUG] 1. File Load Time: {time.time() - t0:.2f} seconds")

    columns_to_delete = set()

    # --- CHECKPOINT 2: Rearranging Columns ---
    t0 = time.time()
    initial_positions = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    
    if 'rawEvent' in initial_positions:
        raw_event_idx = initial_positions['rawEvent']
        sheet.insert_cols(1)
        for row_idx in range(1, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=1, value=sheet.cell(row=row_idx, column=raw_event_idx + 1).value)
        sheet.delete_cols(raw_event_idx + 1)
        
    mid_positions = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}
    
    if 'eventTime' in mid_positions:
        event_time_idx = mid_positions['eventTime']
        sheet.insert_cols(2)
        for row_idx in range(1, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=2, value=sheet.cell(row=row_idx, column=event_time_idx + 1).value)
        sheet.delete_cols(event_time_idx + 1)
    print(f"[DEBUG] 2. Column Rearranging Time: {time.time() - t0:.2f} seconds")

    current_positions = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}
    delete_columns_names = ['rawEventHash', 'aisaacReceivedTime', 'customerURI', 'cenNifiReceiptTime', 'logFilterKafkaInTime', 'logFilterInTime']
    for name in delete_columns_names:
        if name in current_positions:
            columns_to_delete.add(current_positions[name])

    # --- CHECKPOINT 3: Metadata Name Parsing ---
    t0 = time.time()
    with concurrent.futures.ThreadPoolExecutor() as executor:
        futures = [executor.submit(check_name_column, item, sheet) for item in current_positions.items()]
        for future in concurrent.futures.as_completed(futures):
            res = future.result()
            if res:
                col_idx, col_name, unique_value = res
                target_header_name = col_name.replace("Name", "")
                if target_header_name in current_positions:
                    target_col_idx = current_positions[target_header_name]
                    sheet.cell(row=1, column=target_col_idx, value=unique_value)
                    columns_to_delete.add(col_idx)
    print(f"[DEBUG] 3. Metadata Name Check Time: {time.time() - t0:.2f} seconds")

    # --- CHECKPOINT 4: Null Column Scanning ---
    t0 = time.time()
    with concurrent.futures.ThreadPoolExecutor() as executor:
        futures = [executor.submit(check_null_column, col_idx, sheet, columns_to_delete) for col_idx in range(1, sheet.max_column + 1)]
        for future in concurrent.futures.as_completed(futures):
            res = future.result()
            if res is not None:
                columns_to_delete.add(res)
    print(f"[DEBUG] 4. Null Column Scan Time: {time.time() - t0:.2f} seconds")

    # --- CHECKPOINT 5: Deleting Columns ---
    t0 = time.time()
    for col_index in sorted(list(columns_to_delete), reverse=True):
        sheet.delete_cols(col_index)
    print(f"[DEBUG] 5. Column Deletion Time: {time.time() - t0:.2f} seconds")

    # --- CHECKPOINT 6: Styling Headers & Rows ---
    t0 = time.time()
    thin_border = openpyxl.styles.Border(left=openpyxl.styles.Side(style='thin'),
                                       right=openpyxl.styles.Side(style='thin'),
                                       top=openpyxl.styles.Side(style='thin'),
                                       bottom=openpyxl.styles.Side(style='thin'))
    header_font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')

    for cell in sheet[1]:
        cell.border = thin_border
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    row_font = Font(name='Calibri', size=12, color='000000')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.font = row_font
    print(f"[DEBUG] 6. Styling Time: {time.time() - t0:.2f} seconds")

    # --- CHECKPOINT 7: Column Width Calculation ---
    t0 = time.time()
    for column_cells in sheet.columns:
        max_length = 0
        column = openpyxl.utils.get_column_letter(column_cells[0].column)
        col_header = column_cells[0].value
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = (max_length + 2) * 1.2
        if col_header == 'rawEvent':
            sheet.column_dimensions[column].width = 50
        else:
            sheet.column_dimensions[column].width = max(min(adjusted_width, 40), 10)
    print(f"[DEBUG] 7. Width Calculation Time: {time.time() - t0:.2f} seconds")

    # --- CHECKPOINT 8: Saving File ---
    t0 = time.time()
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    base_path, _ = os.path.splitext(filename)
    output_filename = f"{base_path}_{timestamp}.xlsx"
    wb.save(output_filename)
    print(f"[DEBUG] 8. File Save Time: {time.time() - t0:.2f} seconds")

    print(f"\n[SUCCESS] Total Elapsed Time: {time.time() - total_start:.2f} seconds")

if __name__ == "__main__":
    print("=== Log Normalizer (Debug Profiler) ===")
    filename = sys.argv[1] if len(sys.argv) > 1 else input("Enter file path: ").strip().strip('"').strip("'")
    
    threading.Thread(target=input_listener, daemon=True).start()
    
    worker = threading.Thread(target=remove_null_columns, args=(filename,))
    worker.start()
    
    start_wait = time.time()
    warned = False
    while worker.is_alive():
        worker.join(timeout=1.0)
        if time.time() - start_wait > 60 and not warned:
            print("\n[NOTICE] Over 60 seconds elapsed. Enter 'x' to exit: ", end="", flush=True)
            warned = True
    exit_requested.set()