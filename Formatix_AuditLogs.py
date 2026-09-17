import os
import sys
import time
import threading
from datetime import datetime

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter.scrolledtext import ScrolledText


# ============================================================
# APPLICATION STATE
# ============================================================

APP_TITLE = "Normalizer"
APP_BACKGROUND = "#f5f5dc"
APP_ACCENT = "#808080"
APP_SUCCESS = "#2e7d32"
APP_ERROR = "#c62828"
APP_WARNING = "#ef6c00"

processing_active = False
processing_started_at = None
selected_filename = ""
last_output_filename = ""


# ============================================================
# GUI-SAFE HELPERS
# ============================================================

def write_log(message):
    """Add a message to the GUI activity log safely."""
    root.after(0, _write_log_now, message)


def _write_log_now(message):
    log_text.configure(state=tk.NORMAL)
    log_text.insert(tk.END, f"{message}\n")
    log_text.see(tk.END)
    log_text.configure(state=tk.DISABLED)


def set_status(message, color="#333333"):
    """Update the status label safely from any thread."""
    root.after(0, lambda: status_label.config(text=message, fg=color))


def set_stage(message):
    """Update both the status label and activity log."""
    set_status(message, "#333333")
    write_log(message)


def clean_display_path(path):
    """Return a readable absolute path."""
    return os.path.abspath(os.path.expanduser(str(path)))


# ============================================================
# DATA CLEANING FUNCTIONS
# ============================================================

def clean_name_series(series):
    """
    Clean a metadata Name column.

    Invalid values:
    - Empty
    - null
    - nan
    - none
    """
    cleaned_series = series.astype("string").str.strip()
    lowercase_series = cleaned_series.str.lower()

    valid_mask = (
        series.notna()
        & cleaned_series.notna()
        & cleaned_series.ne("")
        & lowercase_series.ne("null")
        & lowercase_series.ne("nan")
        & lowercase_series.ne("none")
    )

    return cleaned_series, valid_mask


def clean_filename_column_name(column_name):
    """
    Clean a dynamic Excel column name.

    Excel column headers cannot be empty. Line breaks and tab
    characters are replaced with spaces.
    """
    column_name = str(column_name).strip()
    column_name = column_name.replace("\r", " ")
    column_name = column_name.replace("\n", " ")
    column_name = column_name.replace("\t", " ")
    return column_name


def get_safe_output_column_name(
    requested_name,
    original_columns,
    generated_columns,
    source_value_column
):
    """Return a safe output column name without overwriting data."""
    requested_name = clean_filename_column_name(requested_name)

    if not requested_name:
        return None

    if requested_name in generated_columns:
        return requested_name

    if requested_name not in original_columns:
        return requested_name

    if requested_name == source_value_column:
        return requested_name

    counter = 2
    safe_name = f"{requested_name}_{counter}"
    all_used_columns = set(original_columns) | set(generated_columns)

    while safe_name in all_used_columns:
        counter += 1
        safe_name = f"{requested_name}_{counter}"

    return safe_name


def expand_metadata_columns_fast(df):
    """
    Convert metadata Name/value pairs into separate columns.

    Important:
    - No rows are added.
    - No rows are deleted.
    - Row order is not changed.
    - Non-applicable cells remain empty.
    - New columns are added once using pd.concat().
    """
    original_index = df.index
    original_row_count = len(df)
    original_columns = list(df.columns)
    original_column_set = set(original_columns)

    metadata_pairs = []

    for name_column in original_columns:
        name_column_text = str(name_column)

        if not name_column_text.endswith("Name"):
            continue

        value_column = name_column_text[:-4]

        if value_column in df.columns:
            metadata_pairs.append((name_column, value_column))

    if not metadata_pairs:
        return df

    generated_columns = {}
    columns_to_drop = []

    for name_column, value_column in metadata_pairs:
        cleaned_names, valid_name_mask = clean_name_series(df[name_column])

        unique_names = (
            cleaned_names.loc[valid_name_mask]
            .drop_duplicates()
            .tolist()
        )

        for unique_name in unique_names:
            requested_name = clean_filename_column_name(unique_name)

            if not requested_name:
                continue

            output_column = get_safe_output_column_name(
                requested_name=requested_name,
                original_columns=original_column_set,
                generated_columns=generated_columns,
                source_value_column=value_column
            )

            if output_column is None:
                continue

            matching_rows = valid_name_mask & cleaned_names.eq(unique_name)

            if output_column not in generated_columns:
                generated_columns[output_column] = pd.Series(
                    pd.NA,
                    index=original_index,
                    dtype="object"
                )

            target_series = generated_columns[output_column]
            source_values = df[value_column]

            target_as_string = (
                target_series
                .astype("string")
                .str.strip()
                .str.lower()
            )

            destination_empty = (
                target_series.isna()
                | target_as_string.isna()
                | target_as_string.eq("")
                | target_as_string.eq("null")
                | target_as_string.eq("nan")
                | target_as_string.eq("none")
            )

            update_mask = matching_rows & destination_empty
            target_series.loc[update_mask] = source_values.loc[update_mask].values
            generated_columns[output_column] = target_series

        columns_to_drop.extend([name_column, value_column])

    columns_to_drop = list(dict.fromkeys(columns_to_drop))
    existing_columns_to_drop = [
        column for column in columns_to_drop if column in df.columns
    ]

    base_df = df.drop(columns=existing_columns_to_drop)

    if generated_columns:
        generated_df = pd.DataFrame(generated_columns, index=original_index)
        result_df = pd.concat([base_df, generated_df], axis=1)
    else:
        result_df = base_df

    result_df = result_df.copy()

    if len(result_df) != original_row_count:
        raise RuntimeError("Metadata expansion changed the number of rows.")

    if not result_df.index.equals(original_index):
        raise RuntimeError("Metadata expansion changed the DataFrame index.")

    return result_df


def drop_empty_columns(df):
    """Remove columns containing only empty/null/nan/none values."""
    columns_to_drop = []

    for column in df.columns:
        cleaned_series = (
            df[column]
            .astype("string")
            .str.strip()
            .str.lower()
        )

        has_valid_data = (
            df[column].notna()
            & cleaned_series.notna()
            & cleaned_series.ne("")
            & cleaned_series.ne("null")
            & cleaned_series.ne("nan")
            & cleaned_series.ne("none")
        )

        if not has_valid_data.any():
            columns_to_drop.append(column)

    if columns_to_drop:
        df = df.drop(columns=columns_to_drop)

    return df


def rearrange_columns(df):
    """Place rawEvent first and eventTime second."""
    columns = list(df.columns)
    front_columns = []

    if "rawEvent" in columns:
        columns.remove("rawEvent")
        front_columns.append("rawEvent")

    if "eventTime" in columns:
        columns.remove("eventTime")
        front_columns.append("eventTime")

    return df[front_columns + columns]


def load_input_file(filename):
    """Load a CSV or XLSX file."""
    extension = os.path.splitext(filename)[1].lower()

    if extension == ".csv":
        return pd.read_csv(filename, low_memory=False)

    if extension == ".xlsx":
        return pd.read_excel(filename, engine="openpyxl")

    raise ValueError("Unsupported file type. Only CSV and XLSX are supported.")


def save_dataframe_to_excel(df, output_filename):
    """Save the cleaned DataFrame to Excel."""
    df.to_excel(output_filename, index=False, engine="openpyxl")


def style_excel_file(output_filename):
    """Apply formatting to the generated Excel workbook."""
    workbook = openpyxl.load_workbook(output_filename)
    sheet = workbook.active
    sheet.title = "Triggered Events"

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    header_font = Font(
        name="Calibri",
        size=13,
        bold=True,
        color="FFFFFF"
    )

    header_fill = PatternFill(
        start_color="4F81BD",
        end_color="4F81BD",
        fill_type="solid"
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    row_font = Font(
        name="Calibri",
        size=12,
        color="000000"
    )

    row_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    for cell in sheet[1]:
        cell.border = thin_border
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.font = row_font
            cell.alignment = row_alignment

    for column_cells in sheet.columns:
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(column_cells[0].column)
        column_header = column_cells[0].value

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        adjusted_width = (max_length + 2) * 1.2

        if column_header == "rawEvent":
            sheet.column_dimensions[column_letter].width = 50
        else:
            sheet.column_dimensions[column_letter].width = max(
                min(adjusted_width, 40),
                10
            )

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = None
    sheet.row_dimensions[1].height = 22
    workbook.save(output_filename)


# ============================================================
# MAIN PROCESSING FUNCTION
# ============================================================

def process_log_file(filename):
    """Normalize the selected file and return processing details."""
    total_start = time.time()

    set_stage("Loading input file...")
    load_start = time.time()
    df = load_input_file(filename)
    load_seconds = time.time() - load_start

    original_row_count = len(df)
    original_column_count = len(df.columns)

    write_log(f"[INFO] File loaded in {load_seconds:.2f} seconds.")
    write_log(f"[INFO] Input rows: {original_row_count}")
    write_log(f"[INFO] Input columns: {original_column_count}")

    set_stage("Rearranging columns...")
    df = rearrange_columns(df)

    set_stage("Removing predefined unwanted columns...")
    unwanted_columns = [
        "rawEventHash",
        "aisaacReceivedTime",
        "customerURI",
        "cenNifiReceiptTime",
        "logFilterKafkaInTime",
        "logFilterInTime"
    ]

    existing_unwanted_columns = [
        column for column in unwanted_columns if column in df.columns
    ]

    if existing_unwanted_columns:
        df = df.drop(columns=existing_unwanted_columns)
        write_log(
            "[INFO] Removed columns: "
            + ", ".join(existing_unwanted_columns)
        )
    else:
        write_log("[INFO] No predefined unwanted columns were found.")

    set_stage("Expanding metadata columns...")
    metadata_start = time.time()
    df = expand_metadata_columns_fast(df)
    metadata_seconds = time.time() - metadata_start
    write_log(
        f"[INFO] Metadata columns expanded in {metadata_seconds:.2f} seconds."
    )

    if len(df) != original_row_count:
        raise RuntimeError("Row count changed after metadata expansion.")

    set_stage("Removing empty columns...")
    null_scan_start = time.time()
    df = drop_empty_columns(df)
    null_scan_seconds = time.time() - null_scan_start
    write_log(
        f"[INFO] Empty columns removed in {null_scan_seconds:.2f} seconds."
    )

    if len(df) != original_row_count:
        raise RuntimeError("Row count changed during processing.")

    set_stage('Converting blank values to "null"...')
    df = df.fillna("null")
    df.replace(
        to_replace=r"^\s*$",
        value="null",
        regex=True,
        inplace=True
    )

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    base_path, _ = os.path.splitext(filename)
    output_filename = f"{base_path}_{timestamp}.xlsx"

    set_stage("Writing Excel file...")
    write_start = time.time()
    save_dataframe_to_excel(df, output_filename)
    write_seconds = time.time() - write_start
    write_log(f"[INFO] Excel data written in {write_seconds:.2f} seconds.")

    set_stage("Applying Excel formatting...")
    style_start = time.time()
    style_excel_file(output_filename)
    style_seconds = time.time() - style_start
    write_log(
        f"[INFO] Excel styling completed in {style_seconds:.2f} seconds."
    )

    total_seconds = time.time() - total_start

    return {
        "output_filename": output_filename,
        "rows": len(df),
        "columns": len(df.columns),
        "elapsed": total_seconds
    }


# ============================================================
# GUI ACTIONS
# ============================================================

def select_file():
    """Open the file picker and start processing the selected file."""
    global selected_filename

    if processing_active:
        messagebox.showwarning(
            "Processing in progress",
            "A file is already being processed."
        )
        return

    filename = filedialog.askopenfilename(
        title="Select CSV or XLSX File",
        filetypes=[
            ("Supported files", "*.csv *.xlsx"),
            ("CSV files", "*.csv"),
            ("Excel files", "*.xlsx"),
            ("All files", "*.*")
        ]
    )

    if not filename:
        return

    filename = clean_display_path(filename)
    extension = os.path.splitext(filename)[1].lower()

    if extension not in (".csv", ".xlsx"):
        messagebox.showerror(
            "Unsupported file",
            "Only CSV and XLSX files are supported."
        )
        return

    if not os.path.isfile(filename):
        messagebox.showerror("File not found", "The selected file was not found.")
        return

    selected_filename = filename
    selected_file_label.config(text=filename)

    log_text.configure(state=tk.NORMAL)
    log_text.delete("1.0", tk.END)
    log_text.configure(state=tk.DISABLED)

    start_processing(filename)


def start_processing(filename):
    """Start file processing on a background thread."""
    global processing_active, processing_started_at

    processing_active = True
    processing_started_at = time.time()

    select_button.config(state=tk.DISABLED)
    open_folder_button.config(state=tk.DISABLED)
    progress_bar.start(12)

    set_status("Processing...", "#333333")
    write_log(f"[INFO] Selected file: {filename}")

    worker = threading.Thread(
        target=processing_worker,
        args=(filename,),
        daemon=True
    )
    worker.start()

    root.after(1000, monitor_processing_time)


def processing_worker(filename):
    """Run data processing without freezing the GUI."""
    try:
        result = process_log_file(filename)
        root.after(0, processing_completed, result)

    except PermissionError:
        error_message = (
            "Permission denied while writing the output file.\n\n"
            "Close the output Excel file if it is already open, "
            "then try again."
        )
        root.after(0, processing_failed, error_message)

    except MemoryError:
        error_message = (
            "The computer does not have enough available memory "
            "to process this file."
        )
        root.after(0, processing_failed, error_message)

    except Exception as error:
        root.after(0, processing_failed, str(error))


def processing_completed(result):
    """Handle successful processing on the main GUI thread."""
    global processing_active, last_output_filename

    processing_active = False
    last_output_filename = result["output_filename"]

    progress_bar.stop()
    select_button.config(state=tk.NORMAL)
    open_folder_button.config(state=tk.NORMAL)

    set_status("Processing completed successfully", APP_SUCCESS)

    write_log("")
    write_log("[SUCCESS] Processing completed.")
    write_log(f"[SUCCESS] Rows preserved: {result['rows']}")
    write_log(f"[SUCCESS] Final columns: {result['columns']}")
    write_log(f"[SUCCESS] Total elapsed time: {result['elapsed']:.2f} seconds")
    write_log(f"[SUCCESS] Saved as: {result['output_filename']}")

    messagebox.showinfo(
        "Processing completed",
        "The file was processed successfully.\n\n"
        f"Rows preserved: {result['rows']}\n"
        f"Final columns: {result['columns']}\n"
        f"Elapsed time: {result['elapsed']:.2f} seconds\n\n"
        f"Saved as:\n{result['output_filename']}"
    )


def processing_failed(error_message):
    """Handle processing failure on the main GUI thread."""
    global processing_active

    processing_active = False
    progress_bar.stop()
    select_button.config(state=tk.NORMAL)

    set_status("Processing failed", APP_ERROR)
    write_log(f"[ERROR] {error_message}")

    messagebox.showerror("Processing failed", error_message)


def monitor_processing_time():
    """Show a notice if processing has run for more than 60 seconds."""
    if not processing_active or processing_started_at is None:
        return

    elapsed = time.time() - processing_started_at

    if elapsed >= 60:
        status_label.config(
            text="Still processing. Large files can take longer.",
            fg=APP_WARNING
        )

    root.after(1000, monitor_processing_time)


def open_output_folder():
    """Open the folder containing the generated Excel file."""
    if not last_output_filename:
        return

    folder = os.path.dirname(os.path.abspath(last_output_filename))

    try:
        if sys.platform.startswith("win"):
            os.startfile(folder)
        elif sys.platform == "darwin":
            import subprocess
            subprocess.Popen(["open", folder])
        else:
            import subprocess
            subprocess.Popen(["xdg-open", folder])
    except Exception as error:
        messagebox.showerror(
            "Unable to open folder",
            f"Could not open the output folder.\n\n{error}"
        )


def exit_application():
    """Close the program, with a warning during active processing."""
    if processing_active:
        should_exit = messagebox.askyesno(
            "Exit Normalizer",
            "A file is currently being processed.\n\n"
            "Exiting now will stop the operation. Do you want to exit?"
        )

        if not should_exit:
            return

    root.destroy()


# ============================================================
# UI CONSTRUCTION
# ============================================================

root = tk.Tk()
root.title(APP_TITLE)
root.geometry("760x590")
root.minsize(700, 540)
root.configure(bg=APP_BACKGROUND)
root.protocol("WM_DELETE_WINDOW", exit_application)

# Optional Windows taskbar icon handling can be added here if an .ico exists.

main_frame = tk.Frame(root, bg=APP_BACKGROUND)
main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)

title_label = tk.Label(
    main_frame,
    text="Normalizer",
    bg=APP_ACCENT,
    fg="white",
    font=("Arial", 24, "bold"),
    pady=12
)
title_label.pack(fill=tk.X)

instruction_label = tk.Label(
    main_frame,
    text="Select a CSV or XLSX file to normalize and export it as a formatted Excel file.",
    bg=APP_BACKGROUND,
    fg="#222222",
    font=("Arial", 11),
    wraplength=680,
    justify=tk.CENTER
)
instruction_label.pack(pady=(18, 8))

selected_file_title = tk.Label(
    main_frame,
    text="Selected file:",
    bg=APP_BACKGROUND,
    fg="#333333",
    font=("Arial", 10, "bold")
)
selected_file_title.pack(anchor="w", pady=(5, 2))

selected_file_label = tk.Label(
    main_frame,
    text="No file selected",
    bg="white",
    fg="#333333",
    font=("Arial", 9),
    relief=tk.SUNKEN,
    anchor="w",
    padx=8,
    pady=7,
    wraplength=680,
    justify=tk.LEFT
)
selected_file_label.pack(fill=tk.X)

button_frame = tk.Frame(main_frame, bg=APP_BACKGROUND)
button_frame.pack(pady=18)

select_button = tk.Button(
    button_frame,
    text="Select File",
    command=select_file,
    bg=APP_ACCENT,
    fg="white",
    activebackground="#696969",
    activeforeground="white",
    font=("Arial", 12, "bold"),
    width=15,
    cursor="hand2"
)
select_button.grid(row=0, column=0, padx=6)

open_folder_button = tk.Button(
    button_frame,
    text="Open Output Folder",
    command=open_output_folder,
    bg=APP_ACCENT,
    fg="white",
    activebackground="#696969",
    activeforeground="white",
    font=("Arial", 12),
    width=18,
    state=tk.DISABLED,
    cursor="hand2"
)
open_folder_button.grid(row=0, column=1, padx=6)

exit_button = tk.Button(
    button_frame,
    text="Exit",
    command=exit_application,
    bg=APP_ACCENT,
    fg="white",
    activebackground="#696969",
    activeforeground="white",
    font=("Arial", 12),
    width=12,
    cursor="hand2"
)
exit_button.grid(row=0, column=2, padx=6)

# ttk is imported here so the normal tkinter imports stay easy to read.
from tkinter import ttk

progress_bar = ttk.Progressbar(
    main_frame,
    mode="indeterminate"
)
progress_bar.pack(fill=tk.X, pady=(0, 8))

status_label = tk.Label(
    main_frame,
    text="Ready",
    bg=APP_BACKGROUND,
    fg=APP_SUCCESS,
    font=("Arial", 11, "bold")
)
status_label.pack(pady=(2, 8))

log_title_label = tk.Label(
    main_frame,
    text="Activity Log",
    bg=APP_BACKGROUND,
    fg="#333333",
    font=("Arial", 10, "bold")
)
log_title_label.pack(anchor="w")

log_text = ScrolledText(
    main_frame,
    height=12,
    font=("Consolas", 9),
    bg="#ffffff",
    fg="#222222",
    state=tk.DISABLED,
    wrap=tk.WORD
)
log_text.pack(fill=tk.BOTH, expand=True, pady=(4, 8))

developer_label = tk.Label(
    main_frame,
    text="Developed by Gourab Sarkar",
    bg=APP_BACKGROUND,
    fg="black",
    font=("Arial", 10)
)
developer_label.pack(anchor="e")


if __name__ == "__main__":
    root.mainloop()
