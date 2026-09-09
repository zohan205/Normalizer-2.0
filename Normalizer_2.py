import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import sys
import os
import time
import tkinter as tk
from tkinter import filedialog, messagebox
import threading

def remove_null_columns(filename):
    start_time = time.time()  # Record start time
    
    # Check the file extension
    ext = os.path.splitext(filename)[1].lower()
    
    try:
        if ext == '.csv':
            # Create a blank in-memory workbook for the CSV data
            wb = openpyxl.Workbook()
            sheet = wb.active
            # Read the CSV and populate the sheet
            with open(filename, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                for r_idx, row in enumerate(reader, 1):
                    for c_idx, val in enumerate(row, 1):
                        sheet.cell(row=r_idx, column=c_idx, value=val)
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(filename)
            sheet = wb.active
        else:
            messagebox.showerror("Error", "Unsupported file type.")
            return
            
    except FileNotFoundError:
        messagebox.showerror("Error", f"The specified file '{filename}' was not found.")
        return
    except Exception as e:
        messagebox.showerror("Error", f"Failed to load file:\n{str(e)}")
        return

    columns_to_delete = set()  # Using a set to prevent duplicate column index markings

    # --- 1. REARRANGE COLUMNS: rawEvent (1st) and eventTime (2nd) ---
    
    # Map initial column layout
    initial_positions = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
    
    # Move 'rawEvent' to the first column
    if 'rawEvent' in initial_positions:
        raw_event_idx = initial_positions['rawEvent']
        sheet.insert_cols(1)
        for row_idx in range(1, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=1, value=sheet.cell(row=row_idx, column=raw_event_idx + 1).value)
        sheet.delete_cols(raw_event_idx + 1)
        
    # Re-map positions before moving eventTime
    mid_positions = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}
    
    # Move 'eventTime' to the second column
    if 'eventTime' in mid_positions:
        event_time_idx = mid_positions['eventTime']
        sheet.insert_cols(2)
        for row_idx in range(1, sheet.max_row + 1):
            sheet.cell(row=row_idx, column=2, value=sheet.cell(row=row_idx, column=event_time_idx + 1).value)
        sheet.delete_cols(event_time_idx + 1)

    # 2. Re-map current positions for processing to avoid index mismatches
    current_positions = {sheet.cell(row=1, column=col).value: col for col in range(1, sheet.max_column + 1)}

    # Mark pre-defined columns for deletion (Removed 'rawEvent' from this list)
    delete_columns_names = ['rawEventHash', 'aisaacReceivedTime', 'customerURI', 'cenNifiReceiptTime', 'logFilterKafkaInTime', 'logFilterInTime']
    for name in delete_columns_names:
        if name in current_positions:
            columns_to_delete.add(current_positions[name])

    # 3. Safely parse and shift mapping titles (e.g., matching "aisaacStr5Name" content onto "aisaacStr5" header)
    for col_name, col_idx in list(current_positions.items()):
        if col_name and str(col_name).endswith("Name"):
            # Check what unique value fills this metadata name column
            unique_values = set()
            for row_idx in range(2, sheet.max_row + 1):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                if cell_value is not None and str(cell_value).strip().lower() != "null":
                    unique_values.add(str(cell_value).strip())

            # If it's a static system descriptor (exactly 1 unique valid name)
            if len(unique_values) == 1:
                unique_value = unique_values.pop()
                target_header_name = col_name.replace("Name", "")
                
                # Relocate lookup value directly into the corresponding variable column's header
                if target_header_name in current_positions:
                    target_col_idx = current_positions[target_header_name]
                    sheet.cell(row=1, column=target_col_idx, value=unique_value)
                    columns_to_delete.add(col_idx)

    # 4. Check for columns that are entirely empty or purely string "null"
    for col_idx in range(1, sheet.max_column + 1):
        if col_idx in columns_to_delete:
            continue
        is_null_column = True
        for row_idx in range(2, sheet.max_row + 1):
            cell_value = sheet.cell(row=row_idx, column=col_idx).value
            if cell_value is not None and str(cell_value).strip().lower() != "null" and str(cell_value).strip() != "":
                is_null_column = False
                break
        if is_null_column:
            columns_to_delete.add(col_idx)

    # 5. Delete flagged columns safely from right to left to protect active indexing
    for col_index in sorted(list(columns_to_delete), reverse=True):
        sheet.delete_cols(col_index)

    # Style formatting for the header row
    for cell in sheet[1]:
        cell.border = openpyxl.styles.Border(left=openpyxl.styles.Side(style='thin'),
                                             right=openpyxl.styles.Side(style='thin'),
                                             top=openpyxl.styles.Side(style='thin'),
                                             bottom=openpyxl.styles.Side(style='thin'))
        cell.font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Style formatting for data rows
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = openpyxl.styles.Border(left=openpyxl.styles.Side(style='thin'),
                                                 right=openpyxl.styles.Side(style='thin'),
                                                 top=openpyxl.styles.Side(style='thin'),
                                                 bottom=openpyxl.styles.Side(style='thin'))
            cell.font = Font(name='Calibri', size=12, color='000000')

    # Adjust column width based on content dynamically (with a cap for JSON objects)
    for column_cells in sheet.columns:
        max_length = 0
        column = openpyxl.utils.get_column_letter(column_cells[0].column)
        col_header = column_cells[0].value
        
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = (max_length + 2) * 1.2
        
        # Capping rawEvent width so it doesn't break the sheet view layout
        if col_header == 'rawEvent':
            sheet.column_dimensions[column].width = 50
        else:
            sheet.column_dimensions[column].width = max(min(adjusted_width, 40), 10)

    # Save the modified workbook as .xlsx regardless of input type
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    base_path, _ = os.path.splitext(filename)
    output_filename = f"{base_path}_{timestamp}.xlsx"
    
    try:
        wb.save(output_filename)
        elapsed_time = time.time() - start_time
        messagebox.showinfo("Success", f"Modified workbook saved as '{output_filename}'.\nElapsed Time: {elapsed_time:.2f} seconds")
    except PermissionError:
        messagebox.showerror("Error", "Permission denied. Please make sure you have write access to the directory.")

def remove_null_columns_threaded(filename):
    threading.Thread(target=remove_null_columns, args=(filename,), daemon=True).start()

def select_file():
    # Allow filtering for both CSV and Excel logs
    file_path = filedialog.askopenfilename(filetypes=[("Excel or CSV Logs", "*.xlsx *.csv")])
    if file_path:
        ext = os.path.splitext(file_path)[1].lower()
        # Enforce strict validation
        if ext not in ['.xlsx', '.csv']:
            messagebox.showerror(
                "Invalid File Type", 
                "Only .csv and .xlsx input files are allowed!\nPlease pick a valid file."
            )
            return
        remove_null_columns_threaded(file_path)

def exit_application():
    sys.exit()

# UI Construction
root = tk.Tk()
root.title("Normalizer")
root.configure(bg='#f5f5dc')

title_label = tk.Label(root, text="Normalizer", bg='#808080', fg='white', font=('Arial', 24), width=30)
title_label.pack(padx=20, pady=(20, 0))

select_button = tk.Button(root, text="Select File", command=select_file, bg='#808080', fg='white', font=('Arial', 14), width=15)
select_button.pack(pady=(40, 10))

exit_button = tk.Button(root, text="Exit", command=exit_application, bg='#808080', fg='white', font=('Arial', 14), width=15)
exit_button.pack(pady=(0, 40))

developer_label = tk.Label(root, text="Developed by Gourab Sarkar", bg='#f5f5dc', fg='black', font=('Arial', 10))
developer_label.pack(side=tk.BOTTOM, anchor='e', padx=10, pady=(0, 10))

root.mainloop()
